from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from mcp.server.fastmcp import FastMCP
from pydantic import BaseModel, Field

from optiagent.mcp_contracts import (
    ProblemEnvelope,
    ProblemSpecModel,
    SourceReference,
    ValidationReport,
)
from optiagent.mcp_servers.common import json_safe, resolve_readable_path, run_server
from optiagent.problem_spec import infer_problem_spec
from optiagent.rag import retrieve


class DocumentContent(BaseModel):
    """文档解析结果。"""

    path: str
    format: str
    text: str
    truncated: bool = False
    metadata: dict[str, Any] = Field(default_factory=dict)
    warnings: list[str] = Field(default_factory=list)


class KnowledgeSearchResult(BaseModel):
    """本地运筹知识库检索结果。"""

    query: str
    documents: list[dict[str, Any]] = Field(default_factory=list)


SUPPORTED_FORMATS = [".pdf", ".docx", ".xlsx", ".xls", ".csv", ".json", ".md", ".txt"]

mcp = FastMCP(
    "OptiAgent Document MCP",
    instructions=(
        "负责读取优化问题相关文档、检索建模知识，并生成不含虚构数据的 ProblemEnvelope。"
        "读取范围受 OPTIAGENT_MCP_ALLOWED_ROOTS 限制。"
    ),
    json_response=True,
)


def read_document_file(path: str, max_chars: int = 30_000, sheet_name: str | None = None) -> DocumentContent:
    """读取支持的本地文档，并转换为可检索文本。"""

    source = resolve_readable_path(path)
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_FORMATS:
        raise ValueError(f"不支持的文件格式：{suffix}")
    if max_chars < 1_000 or max_chars > 200_000:
        raise ValueError("max_chars 必须在 1000 到 200000 之间。")

    warnings: list[str] = []
    metadata: dict[str, Any] = {"size_bytes": source.stat().st_size}
    if suffix in {".md", ".txt"}:
        text = _read_plain_text(source)
    elif suffix == ".csv":
        frame = _read_csv(source)
        text = frame.to_csv(index=False)
        metadata.update({"rows": len(frame), "columns": [str(column) for column in frame.columns]})
    elif suffix == ".json":
        text = _read_plain_text(source)
    elif suffix in {".xlsx", ".xls"}:
        text, excel_meta = _read_excel(source, sheet_name)
        metadata.update(excel_meta)
    elif suffix == ".docx":
        text = _read_docx(source)
    else:
        text, page_count = _read_pdf(source)
        metadata["page_count"] = page_count

    truncated = len(text) > max_chars
    if truncated:
        text = text[:max_chars]
        warnings.append(f"文档内容超过 {max_chars} 字符，已截断。")
    return DocumentContent(
        path=str(source),
        format=suffix.lstrip("."),
        text=text,
        truncated=truncated,
        metadata=json_safe(metadata),
        warnings=warnings,
    )


@mcp.tool(title="读取优化文档")
def document_read(path: str, max_chars: int = 30_000, sheet_name: str | None = None) -> DocumentContent:
    """读取 PDF、Word、Excel、CSV、JSON、Markdown 或文本文件。"""

    return read_document_file(path, max_chars=max_chars, sheet_name=sheet_name)


@mcp.tool(title="检索运筹建模知识")
def document_search_knowledge(query: str, top_k: int = 4) -> KnowledgeSearchResult:
    """从项目本地知识库检索建模、Schema、求解器和代码模板依据。"""

    if not query.strip():
        raise ValueError("检索问题不能为空。")
    top_k = min(max(top_k, 1), 10)
    documents = [
        {
            "title": doc.title,
            "score": round(doc.score, 4),
            "category": doc.category,
            "source": doc.source,
            "content": doc.content,
        }
        for doc in retrieve(query, top_k=top_k)
    ]
    return KnowledgeSearchResult(query=query, documents=documents)


@mcp.tool(title="从文档生成问题结构")
def document_extract_problem(
    path: str,
    question: str = "",
    max_chars: int = 30_000,
    sheet_name: str | None = None,
) -> ProblemEnvelope:
    """读取文档并推断优化问题类型，产生 Solver MCP 可接收的统一外壳。"""

    document = read_document_file(path, max_chars=max_chars, sheet_name=sheet_name)
    context = "\n".join(item for item in [question.strip(), document.text] if item)
    spec = infer_problem_spec(context)
    warnings = [*document.warnings]
    if not question.strip():
        warnings.append("未提供额外问题描述，问题类型仅根据文档内容推断。")
    return ProblemEnvelope(
        problem_spec=ProblemSpecModel.from_problem_spec(spec),
        sources=[SourceReference(kind="file", name=Path(document.path).name, uri=document.path)],
        validation=ValidationReport(
            valid=False,
            warnings=warnings,
            errors=["Document MCP 只生成问题定义；调用 Data MCP 补齐并校验结构化数据。"],
            checks={"document_parsed": True, "data_attached": False},
        ),
        metadata={
            "document": document.metadata,
            "content_excerpt": document.text[:2_000],
        },
    )


def _read_plain_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法识别文件编码：{path.name}")


def _read_csv(path: Path) -> pd.DataFrame:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return pd.read_csv(path, encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法识别 CSV 编码：{path.name}")


def _read_excel(path: Path, sheet_name: str | None) -> tuple[str, dict[str, Any]]:
    try:
        workbook = pd.ExcelFile(path)
    except ImportError as exc:
        raise RuntimeError("读取 Excel 需要安装 openpyxl（.xlsx）或 xlrd（.xls）。") from exc
    selected = sheet_name or workbook.sheet_names[0]
    if selected not in workbook.sheet_names:
        raise ValueError(f"Excel 中不存在工作表：{selected}")
    frame = pd.read_excel(path, sheet_name=selected)
    return frame.to_csv(index=False), {
        "sheet_name": selected,
        "sheet_names": workbook.sheet_names,
        "rows": len(frame),
        "columns": [str(column) for column in frame.columns],
    }


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("读取 Word 需要安装 python-docx。") from exc
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        paragraphs.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(paragraphs)


def _read_pdf(path: Path) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("读取 PDF 需要安装 pypdf。") from exc
    reader = PdfReader(path)
    return "\n\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages)


if __name__ == "__main__":
    run_server(mcp, default_port=8101)
